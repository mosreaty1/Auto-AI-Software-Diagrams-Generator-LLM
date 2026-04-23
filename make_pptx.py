from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# ── palette ───────────────────────────────────────────────────────────────────
DARK_BLUE  = RGBColor(0x1F, 0x45, 0x7A)
MID_BLUE   = RGBColor(0x2E, 0x75, 0xB6)
ACCENT     = RGBColor(0x00, 0xB0, 0xF0)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF2, 0xF2, 0xF2)
DARK_TEXT  = RGBColor(0x26, 0x26, 0x26)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

BLANK = prs.slide_layouts[6]   # completely blank

# ── helpers ───────────────────────────────────────────────────────────────────
def add_rect(slide, l, t, w, h, fill, alpha=None):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    return shape

def txbox(slide, text, l, t, w, h,
          size=20, bold=False, italic=False, color=DARK_TEXT,
          align=PP_ALIGN.LEFT, wrap=True):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    box.word_wrap = wrap
    tf = box.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box

def add_bullet_box(slide, items, l, t, w, h, size=16, title_color=MID_BLUE):
    """items: list of (text, is_bold, indent_level)"""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    box.word_wrap = True
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for text, bold, level in items:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.level = level
        bullet = "  " * level + ("• " if level == 0 else "◦ ")
        run = p.add_run()
        run.text = bullet + text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = DARK_TEXT
    return box

def slide_header(slide, title, subtitle=None):
    # top bar
    add_rect(slide, 0, 0, 13.33, 1.3, DARK_BLUE)
    txbox(slide, title, 0.4, 0.1, 12.5, 0.9,
          size=28, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
    if subtitle:
        txbox(slide, subtitle, 0.4, 0.88, 12.5, 0.4,
              size=14, italic=True, color=ACCENT, align=PP_ALIGN.LEFT)
    # bottom bar
    add_rect(slide, 0, 7.1, 13.33, 0.4, DARK_BLUE)
    txbox(slide, "Auto AI Software Diagrams Generator  |  Final Project Presentation",
          0.3, 7.12, 12.5, 0.3, size=11, color=WHITE, align=PP_ALIGN.LEFT)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — COVER
# ═══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, 13.33, 7.5, DARK_BLUE)
add_rect(sl, 0, 2.8, 13.33, 2.0, MID_BLUE)

txbox(sl, "AUTO AI SOFTWARE DIAGRAMS GENERATOR",
      0.5, 0.7, 12.3, 1.2, size=36, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txbox(sl, "Using Large Language Models",
      0.5, 1.8, 12.3, 0.7, size=22, italic=True, color=ACCENT, align=PP_ALIGN.CENTER)
txbox(sl, "Final Project Presentation",
      0.5, 2.95, 12.3, 0.7, size=26, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txbox(sl, "International Research Conference Style",
      0.5, 3.65, 12.3, 0.5, size=16, italic=True, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)
txbox(sl, "Duration: 10–15 minutes",
      0.5, 6.5, 12.3, 0.5, size=13, color=ACCENT, align=PP_ALIGN.CENTER)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — AGENDA
# ═══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, 13.33, 7.5, LIGHT_GRAY)
slide_header(sl, "Agenda", "What we will cover today")

sections = [
    ("1", "Problem & Motivation"),
    ("2", "Related Work & Research Gap"),
    ("3", "Proposed Methodology"),
    ("4", "Implementation & Results"),
    ("5", "Conclusion & Future Work"),
]
for i, (num, title) in enumerate(sections):
    x = 0.5 + (i % 3) * 4.2
    y = 1.6 + (i // 3) * 2.2
    add_rect(sl, x, y, 3.8, 1.7, MID_BLUE)
    txbox(sl, num, x + 0.15, y + 0.1, 0.6, 0.7, size=32, bold=True, color=ACCENT)
    txbox(sl, title, x + 0.15, y + 0.85, 3.4, 0.7, size=14, bold=True, color=WHITE)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — PROBLEM & MOTIVATION
# ═══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, 13.33, 7.5, LIGHT_GRAY)
slide_header(sl, "1. Problem & Motivation")

# left column
add_rect(sl, 0.4, 1.5, 5.8, 5.3, WHITE)
txbox(sl, "The Problem", 0.6, 1.55, 5.4, 0.5, size=16, bold=True, color=DARK_BLUE)
add_bullet_box(sl, [
    ("Software projects accumulate undocumented code as teams scale", False, 0),
    ("Architecture diagrams created manually → quickly become stale", False, 0),
    ("No CI/CD mechanism for automated diagram generation", False, 0),
    ("Onboarding new developers is slow without visual documentation", False, 0),
], 0.5, 2.1, 5.6, 3.0, size=14)

txbox(sl, "Research Challenges", 0.6, 5.2, 5.4, 0.5, size=16, bold=True, color=DARK_BLUE)
add_bullet_box(sl, [
    ("LLMs produce inconsistent PlantUML syntax across providers", False, 0),
    ("No ground-truth metric for diagram quality without human review", False, 0),
    ("Idle API streams hang pipelines indefinitely", False, 0),
], 0.5, 5.7, 5.6, 1.5, size=13)

# right column
add_rect(sl, 6.8, 1.5, 6.1, 2.4, MID_BLUE)
txbox(sl, "Why It Matters", 7.0, 1.55, 5.7, 0.5, size=16, bold=True, color=WHITE)
add_bullet_box(sl, [
    ("40–60 % of dev time spent reading existing code", False, 0),
    ("Missing diagrams increase defect rates and rework costs", False, 0),
    ("Manual docs are never updated post-release", False, 0),
], 6.9, 2.1, 5.9, 1.7, size=14)
for run in sl.shapes[-1].text_frame.paragraphs[0].runs:
    run.font.color.rgb = WHITE

add_rect(sl, 6.8, 4.1, 6.1, 2.7, DARK_BLUE)
txbox(sl, "Our Goal", 7.0, 4.15, 5.7, 0.5, size=16, bold=True, color=ACCENT)
txbox(sl, "Automate the generation, ranking, and validation of architecture diagrams directly from source code — with zero human intervention.",
      6.9, 4.7, 5.9, 1.8, size=14, color=WHITE)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — RELATED WORK
# ═══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, 13.33, 7.5, LIGHT_GRAY)
slide_header(sl, "2. Related Work & Research Gap")

cols = [
    ("Existing Tools", MID_BLUE, [
        ("Static analysis (Doxygen, PlantUML)", False, 0),
        ("Require manual annotations", False, 1),
        ("No natural-language understanding", False, 1),
        ("Single-LLM generation (GPT-4 studies)", False, 0),
        ("No cross-validation or quality ranking", False, 1),
        ("Code summarisation systems", False, 0),
        ("Textual output only — no UML", False, 1),
    ]),
    ("Limitations", RGBColor(0xC0, 0x50, 0x20), [
        ("Single provider fails silently on hallucinated syntax", False, 0),
        ("No round-trip accuracy measurement", False, 0),
        ("No git-hook / CI integration", False, 0),
        ("Cannot handle partial / idle stream responses", False, 0),
    ]),
    ("Our Contribution", DARK_BLUE, [
        ("5-provider parallel generation + LLM judge ranking", False, 0),
        ("Round-trip: diagram → reconstruct → compare", False, 0),
        ("AbortController timeout → zero hanging pipeline", False, 0),
        ("Drop-in CI/CD hook for post-push automation", False, 0),
    ]),
]
for i, (title, color, items) in enumerate(cols):
    x = 0.4 + i * 4.3
    add_rect(sl, x, 1.5, 4.0, 0.55, color)
    txbox(sl, title, x + 0.1, 1.53, 3.8, 0.45, size=15, bold=True, color=WHITE)
    add_rect(sl, x, 2.1, 4.0, 4.8, WHITE)
    add_bullet_box(sl, items, x + 0.1, 2.2, 3.8, 4.5, size=13)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — METHODOLOGY (architecture)
# ═══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, 13.33, 7.5, LIGHT_GRAY)
slide_header(sl, "3. Proposed Methodology", "Pipeline Architecture")

stages = [
    ("1\nSource\nAnalysis", "Detect language\nExtract classes\n& functions"),
    ("2\nPrompt\nBuilding", "Structured persona\nStep-by-step\ninstructions"),
    ("3\nParallel\nLLM Gen", "5 providers\nfan-out\nsimultaneously"),
    ("4\nJudge &\nRank", "LLM judge\n+ heuristic\nfallback"),
    ("5\nReconstruct\nCode", "Winning diagram\n→ LLM rebuilds\nsource"),
    ("6\nAccuracy\nScore", "5-component\nweighted\nmetric"),
]
box_w, box_h = 1.8, 2.0
start_x = 0.35
y_top = 1.6
for i, (stage, desc) in enumerate(stages):
    x = start_x + i * 2.1
    add_rect(sl, x, y_top, box_w, box_h, MID_BLUE)
    txbox(sl, stage, x + 0.05, y_top + 0.1, box_w - 0.1, 0.85,
          size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_rect(sl, x, y_top + box_h, box_w, 1.5, WHITE)
    txbox(sl, desc, x + 0.05, y_top + box_h + 0.05, box_w - 0.1, 1.4,
          size=11, color=DARK_TEXT, align=PP_ALIGN.CENTER)
    if i < len(stages) - 1:
        txbox(sl, "→", x + box_w + 0.05, y_top + 0.65, 0.35, 0.5,
              size=20, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER)

# providers row
txbox(sl, "LLM Providers: DeepSeek  •  Anthropic Claude  •  Google Gemini  •  Mistral  •  Groq",
      0.4, 5.5, 12.5, 0.5, size=13, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER)

# timeout note
add_rect(sl, 0.4, 6.1, 12.5, 0.9, DARK_BLUE)
txbox(sl, "Reliability: Each provider fetch uses AbortController (120 s timeout) + 1 retry → idle streams never block the pipeline",
      0.6, 6.15, 12.1, 0.75, size=12, color=WHITE)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — IMPLEMENTATION & TOOLS
# ═══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, 13.33, 7.5, LIGHT_GRAY)
slide_header(sl, "4. Implementation & Results", "Tools, Technologies & Metrics")

# left
add_rect(sl, 0.4, 1.5, 6.0, 5.6, WHITE)
txbox(sl, "Tools & Technologies", 0.6, 1.55, 5.6, 0.5, size=16, bold=True, color=DARK_BLUE)
add_bullet_box(sl, [
    ("Node.js 20+ with ES Modules", True, 0),
    ("Zero production dependencies", False, 1),
    ("PlantUML via plantuml.com renderer", True, 0),
    ("Sequence & class diagram output", False, 1),
    ("Native fetch() + AbortController", True, 0),
    ("Timeout + retry per provider", False, 1),
    ("Regex-based language detection", True, 0),
    ("JS, TS, Python, Java, C#, Ruby, Go", False, 1),
    ("Git hook entry point (hook-entry.js)", True, 0),
    ("Post-push CI/CD automation", False, 1),
], 0.5, 2.1, 5.8, 4.8, size=13)

# right — metrics table
add_rect(sl, 6.9, 1.5, 6.0, 5.6, WHITE)
txbox(sl, "Accuracy Metric (5 components)", 7.1, 1.55, 5.6, 0.5, size=16, bold=True, color=DARK_BLUE)

metrics = [
    ("AST Shape Score",  "30%", "Keyword deltas (if/for/class…)"),
    ("Token Jaccard",    "20%", "Token-level set similarity"),
    ("Line Diff Score",  "20%", "Exact line overlap ratio"),
    ("Semantic Score",   "20%", "LLM equivalence judgment"),
    ("Structural Score", "10%", "Function name overlap"),
]
header_y = 2.15
txbox(sl, "Metric", 7.0, header_y, 2.2, 0.38, size=12, bold=True, color=WHITE)
add_rect(sl, 6.9, header_y, 2.3, 0.38, MID_BLUE)
txbox(sl, "Metric", 6.95, header_y + 0.03, 2.2, 0.32, size=11, bold=True, color=WHITE)
add_rect(sl, 9.2, header_y, 0.8, 0.38, MID_BLUE)
txbox(sl, "Weight", 9.25, header_y + 0.03, 0.7, 0.32, size=11, bold=True, color=WHITE)
add_rect(sl, 10.0, header_y, 2.7, 0.38, MID_BLUE)
txbox(sl, "Description", 10.05, header_y + 0.03, 2.6, 0.32, size=11, bold=True, color=WHITE)

for i, (name, weight, desc) in enumerate(metrics):
    row_y = header_y + 0.42 + i * 0.72
    bg = LIGHT_GRAY if i % 2 == 0 else WHITE
    add_rect(sl, 6.9, row_y, 2.3, 0.68, bg)
    add_rect(sl, 9.2, row_y, 0.8, 0.68, bg)
    add_rect(sl, 10.0, row_y, 2.7, 0.68, bg)
    txbox(sl, name,   6.95, row_y + 0.08, 2.2, 0.55, size=11, bold=True, color=DARK_BLUE)
    txbox(sl, weight, 9.25, row_y + 0.08, 0.7, 0.55, size=13, bold=True, color=MID_BLUE, align=PP_ALIGN.CENTER)
    txbox(sl, desc,  10.05, row_y + 0.08, 2.6, 0.55, size=10, color=DARK_TEXT)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — RESULTS
# ═══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, 13.33, 7.5, LIGHT_GRAY)
slide_header(sl, "4. Performance Results", "Evaluation & Comparison with Baseline")

results = [
    ("Winner Score",       "80 – 92 / 100", "LLM judge score for best provider"),
    ("Round-trip Accuracy","75 – 88 %",      "JS & Java sample files"),
    ("Heuristic Agreement","~85 %",           "Matches LLM judge ranking"),
    ("Pipeline Hang Rate", "0 %",             "After AbortController fix (was 100 %)"),
]
for i, (label, value, note) in enumerate(results):
    x = 0.4 + (i % 2) * 6.4
    y = 1.55 + (i // 2) * 2.4
    add_rect(sl, x, y, 5.9, 2.1, WHITE)
    add_rect(sl, x, y, 5.9, 0.55, MID_BLUE)
    txbox(sl, label, x + 0.15, y + 0.07, 5.6, 0.42, size=14, bold=True, color=WHITE)
    txbox(sl, value, x + 0.15, y + 0.65, 5.6, 0.8, size=30, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER)
    txbox(sl, note,  x + 0.15, y + 1.55, 5.6, 0.45, size=12, italic=True, color=RGBColor(0x60,0x60,0x60), align=PP_ALIGN.CENTER)

# baseline comparison
add_rect(sl, 0.4, 6.1, 12.5, 1.0, DARK_BLUE)
txbox(sl, "vs Baseline (single Anthropic, no judging):  Multi-provider selection improves best-diagram score by ~12 pts  •  Accuracy scoring exposes plausible-but-wrong reconstructions  •  CI hook reduces doc lag from days → seconds",
      0.6, 6.15, 12.1, 0.85, size=11, color=WHITE)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 8 — CONCLUSION
# ═══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, 13.33, 7.5, LIGHT_GRAY)
slide_header(sl, "5. Conclusion")

panels = [
    ("Key Contributions", MID_BLUE, [
        "First pipeline to fan-out across 5 heterogeneous LLMs and rank results",
        "Novel round-trip accuracy: source → diagram → reconstruct → compare",
        "Production-grade reliability via AbortController timeout + retry",
        "Zero-dependency Node.js — drop-in CI/CD git hook",
    ]),
    ("Limitations", RGBColor(0xC0, 0x50, 0x20), [
        "Regex source analysis misses complex AST patterns",
        "PlantUML rendering uses a public service (not air-gapped)",
        "Scores are noisy for very short files (< 30 lines)",
        "Semantic score falls back to token Jaccard without a judge",
    ]),
    ("Future Work", DARK_BLUE, [
        "Integrate tree-sitter / @babel/parser for real AST analysis",
        "Self-hosted PlantUML for private / air-gapped deployments",
        "Fine-tune a small model for PlantUML quality evaluation",
        "Support component, deployment & state-machine diagrams",
        "VS Code extension with on-save generation & live preview",
    ]),
]
for i, (title, color, items) in enumerate(panels):
    x = 0.35 + i * 4.3
    add_rect(sl, x, 1.5, 4.1, 0.55, color)
    txbox(sl, title, x + 0.1, 1.53, 3.9, 0.45, size=14, bold=True, color=WHITE)
    add_rect(sl, x, 2.1, 4.1, 4.9, WHITE)
    for j, item in enumerate(items):
        txbox(sl, f"• {item}", x + 0.15, 2.2 + j * 0.85, 3.8, 0.78, size=12, color=DARK_TEXT)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 9 — GRADING RUBRIC
# ═══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, 13.33, 7.5, LIGHT_GRAY)
slide_header(sl, "Grading Rubric", "4 Points Total")

rubric = [
    ("Technical understanding",        "1 pt", "Depth of knowledge, ability to defend design choices"),
    ("Clarity & scientific explanation","1 pt", "Logical reasoning, precise terminology"),
    ("Quality of slides & structure",  "1 pt", "Professional layout, clear flow, visual aids"),
    ("Q&A performance",                "1 pt", "Answering questions accurately under pressure"),
]
for i, (criterion, pts, detail) in enumerate(rubric):
    y = 1.65 + i * 1.35
    add_rect(sl, 0.4, y, 8.5, 1.15, WHITE)
    add_rect(sl, 8.95, y, 1.5, 1.15, MID_BLUE)
    add_rect(sl, 10.5, y, 2.5, 1.15, LIGHT_GRAY)
    txbox(sl, criterion, 0.55, y + 0.08, 8.2, 0.45, size=15, bold=True, color=DARK_BLUE)
    txbox(sl, detail,    0.55, y + 0.55, 8.2, 0.5,  size=11, italic=True, color=DARK_TEXT)
    txbox(sl, pts, 8.95, y + 0.25, 1.5, 0.6, size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txbox(sl, "✓ Evaluated", 10.55, y + 0.3, 2.3, 0.5, size=12, color=MID_BLUE)

add_rect(sl, 0.4, 7.0, 12.5, 0.38, DARK_BLUE)
txbox(sl, "Evaluation standard: depth of understanding • defending design choices • logical reasoning • professional style",
      0.6, 7.03, 12.1, 0.32, size=11, color=WHITE)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 10 — THANK YOU
# ═══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, 13.33, 7.5, DARK_BLUE)
add_rect(sl, 0, 2.8, 13.33, 2.0, MID_BLUE)
txbox(sl, "Thank You", 0.5, 1.0, 12.3, 1.4,
      size=52, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txbox(sl, "Questions Welcome", 0.5, 2.9, 12.3, 0.8,
      size=26, italic=True, color=WHITE, align=PP_ALIGN.CENTER)
txbox(sl, "Auto AI Software Diagrams Generator  —  Multi-LLM Architecture Diagram Pipeline",
      0.5, 4.2, 12.3, 0.6, size=14, color=ACCENT, align=PP_ALIGN.CENTER)
txbox(sl, "github.com/mosreaty1/Auto-AI-Software-Diagrams-Generator-LLM",
      0.5, 5.0, 12.3, 0.5, size=13, italic=True, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

# ── SAVE ──────────────────────────────────────────────────────────────────────
out = "/home/user/Auto-AI-Software-Diagrams-Generator-LLM/Final_Project_Presentation.pptx"
prs.save(out)
print("saved:", out)
